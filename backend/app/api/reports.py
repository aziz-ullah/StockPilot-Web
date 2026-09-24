import io
from datetime import datetime
from decimal import Decimal
from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi import APIRouter, Depends, Query, Response, HTTPException
from sqlalchemy.orm import Session
from sqlalchemy import func
from app.core.database import get_db
from app.api.auth import get_current_user
from app.models.domain import Sale, Purchase, ProductVariant, Expense, Category, User, SaleItem

router = APIRouter(prefix="/reports", tags=["Reports & Exports"])

@router.get("/pnl")
def get_pnl_report(
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    org_id = current_user.organization_id

    sales_q = db.query(
        func.coalesce(func.sum(Sale.total_amount), 0).label("revenue"),
        func.coalesce(func.sum(Sale.total_cogs), 0).label("cogs"),
        func.coalesce(func.sum(Sale.gross_profit), 0).label("gross_profit")
    ).filter(Sale.organization_id == org_id, Sale.status == "Completed")

    exp_q = db.query(
        Expense.category,
        func.coalesce(func.sum(Expense.amount), 0).label("total")
    ).filter(Expense.organization_id == org_id)

    if start_date:
        sales_q = sales_q.filter(Sale.sale_date >= start_date)
        exp_q = exp_q.filter(Expense.expense_date >= start_date)
    if end_date:
        sales_q = sales_q.filter(Sale.sale_date <= end_date)
        exp_q = exp_q.filter(Expense.expense_date <= end_date)

    sales_res = sales_q.one()
    revenue = Decimal(str(sales_res.revenue))
    cogs = Decimal(str(sales_res.cogs))
    gross_profit = Decimal(str(sales_res.gross_profit))

    expense_breakdown = exp_q.group_by(Expense.category).all()
    expense_dict = {cat: Decimal(str(tot)) for cat, tot in expense_breakdown}
    total_expenses = sum(expense_dict.values(), Decimal("0.00"))
    net_profit = gross_profit - total_expenses

    return {
        "start_date": start_date,
        "end_date": end_date,
        "total_revenue": revenue,
        "total_cogs": cogs,
        "gross_profit": gross_profit,
        "expense_breakdown": expense_dict,
        "total_expenses": total_expenses,
        "net_profit": net_profit
    }

@router.get("/export/excel")
def export_reports_to_excel(
    report_type: str = Query(..., enum=["sales", "purchases", "inventory", "pnl", "expenses"]),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = report_type.upper()

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    org_id = current_user.organization_id

    if report_type == "sales":
        headers = ["Invoice #", "Customer Name", "Date", "Payment Method", "Status", "Subtotal ($)", "Discount ($)", "Tax ($)", "Total Revenue ($)", "COGS ($)", "Gross Profit ($)"]
        ws.append(headers)
        sales = db.query(Sale).filter(Sale.organization_id == org_id).order_by(Sale.sale_date.desc()).all()
        for s in sales:
            ws.append([
                s.invoice_number, s.customer_name, s.sale_date.strftime("%Y-%m-%d %H:%M"),
                s.payment_method, s.status, float(s.subtotal), float(s.discount), float(s.tax),
                float(s.total_amount), float(s.total_cogs), float(s.gross_profit)
            ])

    elif report_type == "purchases":
        headers = ["PO Number", "Supplier Name", "Purchase Date", "Status", "Payment Status", "Total Amount ($)", "Notes"]
        ws.append(headers)
        purchases = db.query(Purchase).filter(Purchase.organization_id == org_id).order_by(Purchase.purchase_date.desc()).all()
        for p in purchases:
            ws.append([
                p.purchase_number, p.supplier_name or "N/A", p.purchase_date.strftime("%Y-%m-%d %H:%M"),
                p.status, p.payment_status, float(p.total_amount), p.notes or ""
            ])

    elif report_type == "inventory":
        headers = ["Category", "Product Name", "SKU", "Color", "Size", "Unit Cost ($)", "Selling Price ($)", "Stock Qty", "Min Alert", "Stock Valuation ($)"]
        ws.append(headers)
        variants = db.query(ProductVariant).filter(ProductVariant.organization_id == org_id).all()
        for v in variants:
            p_name = v.product.name if v.product else ""
            c_name = v.product.category.name if v.product and v.product.category else ""
            val = float(v.cost_price * v.stock_quantity)
            ws.append([
                c_name, p_name, v.sku, v.color, v.size, float(v.cost_price), float(v.selling_price),
                v.stock_quantity, v.min_stock_alert, val
            ])

    elif report_type == "expenses":
        headers = ["Category", "Description", "Expense Date", "Amount ($)", "Logged By"]
        ws.append(headers)
        expenses = db.query(Expense).filter(Expense.organization_id == org_id).order_by(Expense.expense_date.desc()).all()
        for e in expenses:
            logged_by = e.created_by.full_name if e.created_by else "N/A"
            ws.append([
                e.category, e.description, e.expense_date.strftime("%Y-%m-%d"), float(e.amount), logged_by
            ])

    elif report_type == "pnl":
        pnl = get_pnl_report(db=db, current_user=current_user)
        ws.append(["PROFIT & LOSS STATEMENT"])
        ws.append([])
        ws.append(["Metric", "Amount ($)"])
        ws.append(["Total Revenue", float(pnl["total_revenue"])])
        ws.append(["Cost of Goods Sold (COGS)", float(pnl["total_cogs"])])
        ws.append(["Gross Profit", float(pnl["gross_profit"])])
        ws.append([])
        ws.append(["Operating Expenses Breakdown"])
        for cat, amt in pnl["expense_breakdown"].items():
            ws.append([f"  - {cat}", float(amt)])
        ws.append(["Total Operating Expenses", float(pnl["total_expenses"])])
        ws.append([])
        ws.append(["Net Profit (Gross Profit - Expenses)", float(pnl["net_profit"])])

    # Styling header
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=StockPilot_{report_type.title()}_Report.xlsx"}
    )
